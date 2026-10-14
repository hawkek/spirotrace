"""Blind coding workbooks: one per coder, one sheet per report.

A sheet shows the record id, the layout, start and end time cells and one
row per core field of that layout, with empty value, value state, page and
note cells. The value state cell is a dropdown of the five ``ValueState``
members. Nothing else is written: the function receives record ids, layout
labels and field labels only, never a value from any run.

Value cells are text-formatted so Excel keeps what the coder types ("0.75",
"75", a leading zero, a date) instead of reinterpreting it.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..model import ColumnRole, Parameter, ValueState
from .allocation import Allocation
from .frame import EvaluateError
from .sampling import stage_rng

WORKBOOK_NAME = "coding_{coder}_CONTAINS_DATA.xlsx"
"""Workbooks hold record ids (participant numbers), hence the data label."""

INSTRUCTIONS_SHEET = "Instructions"
HEADERS = ("field", "value", "value_state", "page", "note")
FIRST_FIELD_ROW = 7
MAX_SHEET_TITLE = 31  # Excel's limit

VALUE_STATE_MEANINGS: Mapping[ValueState, str] = {
    ValueState.VALUE: "a value is printed: type it exactly as printed",
    ValueState.EMPTY: "the cell is blank",
    ValueState.DASH: "a dash or similar mark is printed instead of a value",
    ValueState.NOT_ON_PAGE: "this field or column does not exist on this report",
    ValueState.UNREADABLE: "something is printed but it cannot be read",
}

DEFAULT_INSTRUCTIONS = (
    "Code each sheet from the report itself. Type values exactly as printed, including "
    "any sign, decimal mark or percent sign. Choose a value state on every row. Give the "
    "page number the value was read from. Enter the start time before the first row and "
    "the end time after the last. Do not look at any extracted data while coding."
)


@dataclass(frozen=True)
class CoreField:
    """One field coders record: a parameter in one column role, with a label for the row."""

    parameter: Parameter
    role: ColumnRole
    label: str

    def __post_init__(self) -> None:
        if not isinstance(self.parameter, Parameter):
            raise EvaluateError(
                f"core field parameter must be a Parameter, got {self.parameter!r}"
            )
        if not isinstance(self.role, ColumnRole):
            raise EvaluateError(
                f"core field role must be a ColumnRole, got {self.role!r}"
            )
        if not self.label.strip():
            raise EvaluateError("core field label is blank")


CoreFieldSpec = Mapping[str, tuple[CoreField, ...]]
"""Core fields per layout label."""


def check_core_fields(core_fields: CoreFieldSpec, layouts: Iterable[str]) -> None:
    """Every layout has fields, with no parameter and role listed twice."""
    for layout in layouts:
        fields = core_fields.get(layout)
        if not fields:
            raise EvaluateError(f"no core fields for layout {layout!r}")
        keys = [(f.parameter, f.role) for f in fields]
        if len(set(keys)) != len(keys):
            raise EvaluateError(f"a core field is listed twice for layout {layout!r}")
        labels = [f.label for f in fields]
        if len(set(labels)) != len(labels):
            raise EvaluateError(
                f"a core field label is used twice for layout {layout!r}"
            )


def value_state_list() -> str:
    """The dropdown formula: the five value states, in vocabulary order."""
    return '"' + ",".join(state.value for state in ValueState) + '"'


def workbook_paths(out_dir: Path, coders: Iterable[str]) -> dict[str, Path]:
    return {coder: out_dir / WORKBOOK_NAME.format(coder=coder) for coder in coders}


def write_coder_workbooks(
    allocation: Allocation,
    core_fields: CoreFieldSpec,
    out_dir: Path,
    seed: int,
    instructions: str = DEFAULT_INSTRUCTIONS,
) -> dict[str, Path]:
    """Write one workbook per coder; return the paths by coder label.

    Sheets are in an order drawn from the seed and the coder's position, so
    the two coders of an overlap report meet it at different points. Refuses
    to overwrite: a coder may already have typed into an existing workbook.
    """
    check_core_fields(core_fields, allocation.layouts)
    paths = workbook_paths(out_dir, allocation.coders)
    existing = [p.name for p in paths.values() if p.exists()]
    if existing:
        raise EvaluateError(f"refusing to overwrite existing workbooks: {existing}")
    out_dir.mkdir(parents=True, exist_ok=True)
    for index, coder in enumerate(allocation.coders):
        assignments = {a.record_id: a for a in allocation.for_coder(coder)}
        order = sorted(assignments)
        stage_rng(seed, f"order:{index}").shuffle(order)
        workbook = Workbook()
        _write_instructions(workbook.worksheets[0], coder, len(order), instructions)
        for record_id in order:
            if len(record_id) > MAX_SHEET_TITLE:
                raise EvaluateError("a record id is too long for an Excel sheet title")
            sheet = workbook.create_sheet(record_id)
            layout = assignments[record_id].layout
            _write_report_sheet(sheet, record_id, layout, core_fields[layout])
        workbook.save(paths[coder])
    return paths


def _write_instructions(
    sheet: Worksheet, coder: str, n_reports: int, instructions: str
) -> None:
    sheet.title = INSTRUCTIONS_SHEET
    bold = Font(bold=True)
    sheet["A1"], sheet["B1"] = "Coder", coder
    sheet["A2"], sheet["B2"] = "Reports", n_reports
    sheet["A4"], sheet["B4"] = "Value state", "Meaning"
    sheet["A1"].font = sheet["A2"].font = sheet["A4"].font = sheet["B4"].font = bold
    row = 5
    for state, meaning in VALUE_STATE_MEANINGS.items():
        sheet.cell(row=row, column=1, value=state.value)
        sheet.cell(row=row, column=2, value=meaning)
        row += 1
    note = sheet.cell(row=row + 1, column=1, value=instructions)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=2)
    sheet.row_dimensions[row + 1].height = 110
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 70


def _write_report_sheet(
    sheet: Worksheet, record_id: str, layout: str, fields: tuple[CoreField, ...]
) -> None:
    bold = Font(bold=True)
    for row, (label, value) in enumerate(
        (
            ("Record", record_id),
            ("Layout", layout),
            ("Start time", None),
            ("End time", None),
        ),
        start=1,
    ):
        sheet.cell(row=row, column=1, value=label).font = bold
        cell = sheet.cell(row=row, column=2, value=value)
        cell.number_format = "@"
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=FIRST_FIELD_ROW - 1, column=column, value=header).font = bold
    last_row = FIRST_FIELD_ROW + len(fields) - 1
    for row, core_field in enumerate(fields, start=FIRST_FIELD_ROW):
        sheet.cell(row=row, column=1, value=core_field.label)
        sheet.cell(row=row, column=2).number_format = "@"
    dropdown = DataValidation(
        type="list", formula1=value_state_list(), allow_blank=True
    )
    dropdown.error = "Choose one of the listed value states."
    dropdown.errorTitle = "Value state"
    sheet.add_data_validation(dropdown)
    dropdown.add(f"C{FIRST_FIELD_ROW}:C{last_row}")
    for column, width in zip(
        range(1, len(HEADERS) + 1), (26, 14, 16, 8, 40), strict=True
    ):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{FIRST_FIELD_ROW}"
