"""The dataset views of the resolved values (stage 9): values long, dataset wide,
the status file, the free-text file and the workbook.

Every view is generated from ``resolved`` and ``derived`` and never edited. A
wide cell is blank when the value is not usable (unreadable, excluded, under
reviewer disagreement, or carrying an open error), or when the cell holds no
value (empty, dash, not on the page); the reason goes to the status file.
Parameters marked as free text or as site identifiers never enter the wide
table: they go to the free-text file only.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from spirotrace.derive import DERIVED_NAMES, DerivedRow, DeriveRun
from spirotrace.model import (
    FREE_TEXT_PARAMETERS,
    RATIO_PARAMETERS,
    SITE_IDENTIFIER_PARAMETERS,
    TRIAL_ROLES,
    VALUE_ROLES,
    ColumnRole,
    ObservationKey,
    Parameter,
    ParameterKind,
    ValueState,
    kind_of,
)
from spirotrace.resolve import Resolved, ResolvedValue
from spirotrace.store import Notice, RawStore, Row, write_bytes, write_rows

VALUES_LONG_FILE = "values_long_CONTAINS_DATA.csv"
DATASET_WIDE_FILE = "dataset_wide_CONTAINS_DATA.csv"
STATUS_LONG_FILE = "status_long_CONTAINS_DATA.csv"
FREE_TEXT_FILE = "free_text_CONTAINS_DATA.csv"
WORKBOOK_FILE = "dataset_CONTAINS_DATA.xlsx"

TREND_ROLES: frozenset[ColumnRole] = frozenset(
    {ColumnRole.TREND, ColumnRole.TREND_PCTPRED, ColumnRole.TREND_Z}
)
"""Trend values have one row per visit position: values long only, never the wide table."""

OUT_OF_DATASET: frozenset[Parameter] = FREE_TEXT_PARAMETERS | SITE_IDENTIFIER_PARAMETERS
"""Free text and site identifiers (device serial, software version): the free-text file only."""

FIXED_COLUMNS: tuple[str, ...] = (
    "record_id",
    "layout",
    "pages",
    "n_open_errors",
    "n_open_checks",
    "n_corrected",
)
VALUES_LONG_COLUMNS: tuple[str, ...] = (
    "record_id",
    "layout",
    "name",
    "value",
    "value_state",
    "status",
    "struck",
    "raw_text",
    "read_method",
    "page",
    "machine_confirmed_by",
)
STATUS_LONG_COLUMNS: tuple[str, ...] = (
    "record_id",
    "name",
    "value_state",
    "status",
    "open_severity",
    "reason",
)
FREE_TEXT_COLUMNS: tuple[str, ...] = (
    "record_id",
    "page",
    "parameter",
    "column_role",
    "raw_text",
    "parsed_text",
    "status",
)

_PARAMETER_ORDER = {p: i for i, p in enumerate(Parameter)}
_ROLE_ORDER = {r: i for i, r in enumerate(ColumnRole)}


def _no_notice(_: str) -> None:
    return None


def column_name(key: ObservationKey) -> str:
    """``<parameter>_<role>``, with ``_<position>`` for trial and trend roles."""
    name = f"{key.parameter.value}_{key.column_role.value}"
    if key.column_role in TRIAL_ROLES or key.column_role in TREND_ROLES:
        name += f"_{key.display_position}"
    return name


def column_sort_key(key: ObservationKey) -> tuple[int, int, int]:
    return (
        _PARAMETER_ORDER[key.parameter],
        _ROLE_ORDER[key.column_role],
        key.display_position,
    )


def in_wide_table(key: ObservationKey) -> bool:
    return key.parameter not in OUT_OF_DATASET and key.column_role not in TREND_ROLES


def format_value(rv: ResolvedValue) -> str:
    """The resolved value as the dataset prints it: ratios as three-decimal fractions,
    other numbers as read (no trailing zeros added or lost beyond the float), text as is."""
    r = rv.resolved
    if r.value_state is not ValueState.VALUE:
        return ""
    if r.parsed_value is not None:
        if rv.key.parameter in RATIO_PARAMETERS and rv.key.column_role in VALUE_ROLES:
            return f"{r.parsed_value:.3f}"
        return f"{r.parsed_value:g}"
    return r.parsed_text or ""


def blank_reason(rv: ResolvedValue) -> str:
    """Why a wide cell is blank, or ``""`` when it is not."""
    if rv.status in ("unreadable", "excluded", "reviewer_disagreement"):
        return rv.status
    if rv.open_severity == "error":
        return "open error"
    if rv.resolved.value_state is not ValueState.VALUE:
        return rv.resolved.value_state.value
    return ""


def _first_page(
    values: Mapping[ObservationKey, ResolvedValue],
) -> dict[tuple[str, str], ResolvedValue]:
    """One value per (record, column): the first page's reading when a header field was
    read on two pages (decision 3: both kept as observations, the dataset takes page 1)."""
    out: dict[tuple[str, str], ResolvedValue] = {}
    for key, rv in sorted(
        values.items(), key=lambda kv: (kv[0].page, column_sort_key(kv[0]))
    ):
        slot = (rv.observation.record_id, column_name(key))
        out.setdefault(slot, rv)
    return out


@dataclass(slots=True)
class Dataset:
    """The generated views, in memory, ready to write."""

    wide_columns: list[str]
    wide_rows: list[Row]
    values_long: list[Row]
    status_long: list[Row]
    free_text: list[Row]
    column_keys: dict[str, ObservationKey] = field(default_factory=dict)
    """A representative key per wide column (for the dictionary's kind and unit)."""
    units: dict[str, str] = field(default_factory=dict)
    derived_columns: list[str] = field(default_factory=list)
    blank_counts: Counter[str] = field(default_factory=Counter)


def build_dataset(store: RawStore, resolved: Resolved, derived: DeriveRun) -> Dataset:
    """Build every view from the resolved values and the derived rows."""
    layout_of = {r.record_id: r.record.layout for r in store.records}
    pages_of = {r.record_id: r.record.pages for r in store.records}
    first = _first_page(resolved.values)
    columns: dict[str, ObservationKey] = {}
    units: dict[str, str] = {}
    for (_rid, name), rv in first.items():
        if in_wide_table(rv.key):
            columns.setdefault(name, rv.key)
            if rv.observation.unit and name not in units:
                units[name] = rv.observation.unit
    wide_names = sorted(columns, key=lambda n: column_sort_key(columns[n]))
    derived_names = list(DERIVED_NAMES)
    gli_names = sorted({r.name for r in derived.rows if r.name.startswith("gli_")})
    derived_names = [n for n in derived_names if any(r.name == n for r in derived.rows)]
    derived_names += gli_names
    by_record_derived: dict[str, dict[str, DerivedRow]] = {}
    for r in derived.rows:
        by_record_derived.setdefault(r.record_id, {})[r.name] = r

    wide_rows: list[Row] = []
    status_rows: list[Row] = []
    long_rows: list[Row] = []
    free_rows: list[Row] = []
    blank_counts: Counter[str] = Counter()
    records = sorted(store.records, key=lambda r: r.record_id)
    per_record: dict[str, dict[str, ResolvedValue]] = {}
    for (rid, name), rv in first.items():
        per_record.setdefault(rid, {})[name] = rv

    for rec in records:
        rid = rec.record_id
        values = per_record.get(rid, {})
        all_values = resolved.of_record(rid)
        n_err = sum(1 for v in all_values.values() if v.open_severity == "error")
        n_chk = sum(1 for v in all_values.values() if v.open_severity == "check")
        n_cor = sum(1 for v in all_values.values() if v.status == "corrected")
        row: Row = {
            "record_id": rid,
            "layout": layout_of.get(rid, ""),
            "pages": " ".join(str(p) for p in pages_of.get(rid, ())),
            "n_open_errors": str(n_err),
            "n_open_checks": str(n_chk),
            "n_corrected": str(n_cor),
        }
        for name in wide_names:
            got = values.get(name)
            if got is None:
                row[name] = ""
                continue
            reason = blank_reason(got)
            if reason:
                row[name] = ""
                blank_counts[name] += 1
                status_rows.append(
                    {
                        "record_id": rid,
                        "name": name,
                        "value_state": got.resolved.value_state.value,
                        "status": got.status,
                        "open_severity": got.open_severity,
                        "reason": reason,
                    }
                )
            else:
                row[name] = format_value(got)
        drows = by_record_derived.get(rid, {})
        for name in derived_names:
            d = drows.get(name)
            row[name] = "" if d is None else d.value
            if d is not None and d.value == "":
                status_rows.append(
                    {
                        "record_id": rid,
                        "name": name,
                        "value_state": "",
                        "status": d.status,
                        "open_severity": "",
                        "reason": d.status,
                    }
                )
        wide_rows.append(row)
        # the long views take every observation, both pages, trend rows included
        for key, rv in sorted(
            all_values.items(), key=lambda kv: (column_sort_key(kv[0]), kv[0].page)
        ):
            if key.parameter in OUT_OF_DATASET:
                free_rows.append(
                    {
                        "record_id": rid,
                        "page": str(key.page),
                        "parameter": key.parameter.value,
                        "column_role": key.column_role.value,
                        "raw_text": rv.resolved.raw_text,
                        "parsed_text": rv.resolved.parsed_text or "",
                        "status": rv.status,
                    }
                )
                continue
            reason = blank_reason(rv)
            long_rows.append(
                {
                    "record_id": rid,
                    "layout": layout_of.get(rid, ""),
                    "name": column_name(key),
                    "value": "" if reason else format_value(rv),
                    "value_state": rv.resolved.value_state.value,
                    "status": rv.status,
                    "struck": "true" if rv.resolved.struck else "false",
                    "raw_text": rv.resolved.raw_text,
                    "read_method": rv.resolved.read_method.value,
                    "page": str(key.page),
                    "machine_confirmed_by": " ".join(rv.machine_confirmed_by),
                }
            )
    return Dataset(
        [*FIXED_COLUMNS, *wide_names, *derived_names],
        wide_rows,
        long_rows,
        status_rows,
        free_rows,
        columns,
        units,
        derived_names,
        blank_counts,
    )


def kind_label(key: ObservationKey) -> str:
    return kind_of(key.parameter).value


def is_numeric_column(key: ObservationKey) -> bool:
    return kind_of(key.parameter) is ParameterKind.NUMBER


def write_workbook(
    path: Path,
    dataset: Dataset,
    dictionary_rows: Sequence[Row],
    dictionary_columns: Sequence[str],
    derived_rows: Iterable[DerivedRow],
    notice: Notice = _no_notice,
) -> Path:
    """``dataset_CONTAINS_DATA.xlsx``: Data, Status, Dictionary and Derived sheets.

    Plain cells, no colours or comments, so nothing invites hand editing; numbers
    are written as numbers where the column is numeric.
    """
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(
        title: str,
        columns: Sequence[str],
        rows: Iterable[Mapping[str, str]],
        numeric: set[str],
    ):
        ws = wb.create_sheet(title)
        ws.append(list(columns))
        for c in ws[1]:
            c.font = bold
        for r in rows:
            out: list[object] = []
            for c in columns:
                v = r.get(c, "")
                if c in numeric and v != "":
                    try:
                        out.append(float(v))
                        continue
                    except ValueError:
                        pass
                out.append(v)
            ws.append(out)
        for i, c in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(
                10, min(40, len(c) + 2)
            )
        ws.freeze_panes = "B2"
        return ws

    numeric = {n for n, k in dataset.column_keys.items() if is_numeric_column(k)} | {
        "n_open_errors",
        "n_open_checks",
        "n_corrected",
    }
    numeric |= {
        n
        for n in dataset.derived_columns
        if n.startswith(("gli_", "bdr_")) and not n.startswith("bdr_positive")
    }
    wb.remove(wb.worksheets[0])
    sheet("Data", dataset.wide_columns, dataset.wide_rows, numeric)
    sheet("Status", STATUS_LONG_COLUMNS, dataset.status_long, set())
    sheet("Dictionary", dictionary_columns, dictionary_rows, {"n_records", "n_blank"})
    sheet("Derived", DerivedRow.COLUMNS, (r.to_row() for r in derived_rows), set())
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return write_bytes(path, buf.getvalue(), notice)


def write_dataset_files(
    folder: Path, dataset: Dataset, notice: Notice = _no_notice
) -> list[Path]:
    """The four CSV views; returns the paths written."""
    return [
        write_rows(
            folder / VALUES_LONG_FILE, VALUES_LONG_COLUMNS, dataset.values_long, notice
        ),
        write_rows(
            folder / DATASET_WIDE_FILE, dataset.wide_columns, dataset.wide_rows, notice
        ),
        write_rows(
            folder / STATUS_LONG_FILE, STATUS_LONG_COLUMNS, dataset.status_long, notice
        ),
        write_rows(
            folder / FREE_TEXT_FILE, FREE_TEXT_COLUMNS, dataset.free_text, notice
        ),
    ]
