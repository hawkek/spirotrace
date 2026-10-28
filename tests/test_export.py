"""export: column names, the wide table and its blanking rule, the long and status
files, the free-text file, the dictionary, the QC report tables and HTML, and the
end-to-end export through the pipeline with the share/ files guarded."""

from __future__ import annotations

from pathlib import Path

from checks_helpers import P, R, cells, run
from openpyxl import load_workbook

from spirotrace.config import default_settings
from spirotrace.derive import derive_records
from spirotrace.export import (
    DATASET_WIDE_FILE,
    DICTIONARY_FILE,
    FREE_TEXT_FILE,
    QC_REPORT_FILE,
    STATUS_LONG_FILE,
    VALUES_LONG_FILE,
    WORKBOOK_FILE,
    build_dataset,
    build_tables,
    column_name,
    dictionary_rows,
    export_dataset,
    render_html,
    strip_box_svg,
    summary_lines,
)
from spirotrace.model import ObservationKey, Page, TextSource
from spirotrace.privacy import MASK, QC_SUMMARY_FILE, VARIABLE_SUMMARY_FILE
from spirotrace.resolve import merge_events, resolve_records
from spirotrace.store import Event, RawStore, ReadRecord, read_rows

HASH = "ab" * 32
LAYOUT = "layout_x"

SPEC = {
    (P.FEV1, R.PRE): "3.05",
    (P.FVC, R.PRE): "3.70",
    (P.FEV1_FVC, R.PRE): "82",
    (P.FEV1, R.Z): "-0.50",
    (P.FVC, R.Z): "x",
    (P.FEV1_FVC, R.Z): "-0.10",
    (P.FEV1, R.TRIAL, 1): "3.05",
    (P.FEV1, R.TRIAL, 2): "3.05",
    (P.FEV1, R.TRIAL, 3): "2.98",
    (P.FEV1, R.TREND): "3.00",
    (P.AGE, R.MEASURED): "40",
    (P.SEX, R.MEASURED): "Female",
    (P.TECHNICAL_COMMENTS, R.MEASURED): "Patient coughed on trial 2.",
    (P.DEVICE_SERIAL, R.MEASURED): "SN-1234",
    (P.AGE, R.MEASURED, 1, 2): "40",  # the same header field on page 2 (decision 3)
}


def key(p, r, pos=1, page=1):
    return ObservationKey(HASH, page, LAYOUT, p, r, pos)


def ev(seq, action, value="", k=None, reviewer="KH"):
    return Event(
        f"{seq:032x}",
        seq,
        f"2026-11-02T10:{seq:02d}:00+13:00",
        reviewer,
        "11",
        (k or key(P.FVC, R.Z)).to_string(),
        action,
        value,
        shown_at=f"2026-11-02T10:{seq:02d}:00+13:00",
        elapsed_ms=4000,
        note="n" if action == "exclude" else "",
    )


def pipeline_objects(events=()):
    rec = cells(SPEC)
    store = RawStore(
        records=[ReadRecord(rec.record, dict(rec.observations), list(rec.conditions))],
        pages=[Page(HASH, 1, "summary", 1, 40, 0, 0.0, 0.2, TextSource.TEXT_LAYER)],
        flags=run(rec),
        file_names={HASH: "11.pdf"},
    )
    s = default_settings()
    resolved = resolve_records(store, merge_events(list(events)), s)
    derived = derive_records(resolved, s)
    return store, resolved, derived


# --- naming and the wide table -----------------------------------------------------


def test_column_names() -> None:
    assert column_name(key(P.FEV1, R.PRE)) == "FEV1_pre"
    assert column_name(key(P.FEV1, R.POST_PCTPRED)) == "FEV1_post_pctpred"
    assert column_name(key(P.FEV1, R.TRIAL, 2)) == "FEV1_trial_2"
    assert column_name(key(P.FEV1, R.POST_TRIAL, 3)) == "FEV1_post_trial_3"
    assert column_name(key(P.TLC, R.MEASURED)) == "TLC_measured"
    assert column_name(key(P.AGE, R.MEASURED)) == "age_measured"


def test_wide_table_blanks_by_the_derive_rule_with_reasons() -> None:
    store, resolved, derived = pipeline_objects()
    d = build_dataset(store, resolved, derived)
    row = d.wide_rows[0]
    assert (
        row["record_id"] == "11" and row["layout"] == LAYOUT and row["pages"] == "1 2"
    )
    assert row["FEV1_pre"] == "3.05" and row["FEV1_FVC_pre"] == "0.820"
    assert row["FVC_z"] == ""  # not a number: open error
    assert row["age_measured"] == "40"  # page 1 wins over page 2
    assert "FEV1_trend_1" not in row and "FEV1_trial_2" in row
    assert (
        "technical_comments_measured" not in row and "device_serial_measured" not in row
    )
    assert row["pattern"] == "" and row["n_open_errors"] == "1"
    reasons = {(r["name"], r["reason"]) for r in d.status_long}
    assert ("FVC_z", "open error") in reasons and ("pattern", "not computed") in reasons
    assert d.wide_columns[:6] == [
        "record_id",
        "layout",
        "pages",
        "n_open_errors",
        "n_open_checks",
        "n_corrected",
    ]
    assert (
        d.wide_columns.index("FEV1_pre")
        < d.wide_columns.index("FVC_pre")
        < d.wide_columns.index("age_measured")
    )


def test_correction_fills_the_cell_and_derives() -> None:
    store, resolved, derived = pipeline_objects([ev(1, "typed", "-0.3")])
    d = build_dataset(store, resolved, derived)
    row = d.wide_rows[0]
    assert (
        row["FVC_z"] == "-0.3"
        and row["pattern"] == "normal"
        and row["n_corrected"] == "1"
    )


def test_long_free_text_and_site_identifiers() -> None:
    store, resolved, derived = pipeline_objects()
    d = build_dataset(store, resolved, derived)
    names = {r["name"] for r in d.values_long}
    assert "FEV1_trend_1" in names and "age_measured" in names
    assert not any(
        r["name"].startswith(("technical_comments", "device_serial"))
        for r in d.values_long
    )
    ft = {r["parameter"]: r["raw_text"] for r in d.free_text}
    assert (
        ft["technical_comments"].startswith("Patient")
        and ft["device_serial"] == "SN-1234"
    )
    both_pages = [r for r in d.values_long if r["name"] == "age_measured"]
    assert sorted(r["page"] for r in both_pages) == ["1", "2"]


# --- dictionary ----------------------------------------------------------------------


def test_dictionary_describes_every_column_and_every_rule() -> None:
    store, resolved, derived = pipeline_objects()
    d = build_dataset(store, resolved, derived)
    rows = dictionary_rows(d, 1)
    by_name = {(r["section"], r["name"]): r for r in rows}
    assert set(d.wide_columns) <= {n for s, n in by_name if s == "column"}
    trial = by_name[("column", "FEV1_trial_2")]
    assert (
        "position 1 is the Best column" in trial["description"]
        and trial["display_position"] == "2"
    )
    ratio = by_name[("column", "FEV1_FVC_pre")]
    assert ratio["stored_form"] == "fraction" and ratio["kind"] == "number"
    assert by_name[("column", "FVC_z")]["n_blank"] == "1"
    assert by_name[("column", "pattern")]["source"] == "derived"
    assert ("rule", "range.outside_possible") in by_name
    assert by_name[("rule", "range.outside_possible")]["kind"] == "range"
    assert all(r["description"] for r in rows)


# --- QC report ----------------------------------------------------------------------


def test_qc_tables_and_html(tmp_path: Path) -> None:
    events = [ev(1, "typed", "-0.3"), ev(2, "as_read", "", key(P.FEV1, R.PRE))]
    store, resolved, derived = pipeline_objects(events)
    d = build_dataset(store, resolved, derived)
    sessions = [
        {
            "session_id": "s",
            "reviewer": "KH",
            "started": "2026-11-02T10:00:00+13:00",
            "ended": "2026-11-02T10:30:00+13:00",
            "answers": "2",
        }
    ]
    t = build_tables(store, resolved, derived, d, events, sessions, "", "")
    run_ = dict(t.run)
    assert run_["records"] == "1" and run_["records: layout_x"] == "1"
    by_rule = {r["rule_id"]: r for r in t.flags_by_rule[1]}
    assert (
        by_rule["ocr.not_a_number"]["answered"] == "1"
        and by_rule["ocr.not_a_number"]["value_changed"] == "1"
    )
    assert by_rule["ocr.not_a_number"]["cleared"] == "1"
    effort = dict(t.effort)
    assert (
        effort["reviewer 1: answers"] == "2"
        and effort["reviewer 1: wall-clock minutes (sessions)"] == "30.0"
    )
    assert "KH" not in str(t.effort)
    comp = {(r["variable"], r["layout"]): r for r in t.completeness[1]}
    assert comp[("FVC_z", LAYOUT)]["n_corrected"] == "1"
    dist = {r["name"]: r for r in t.distributions[1]}
    assert dist["FEV1_pre"]["n"] == "1" and dist["FEV1_pre"]["median"] == "3.05"
    html = render_html(t)
    assert "<svg" in html and "FEV1_pre" in html and "KH" not in html
    assert "3.05" in html  # values appear: the file is named _CONTAINS_DATA
    lines = summary_lines(t)
    assert any(line.startswith("records: 1") for line in lines)
    assert not any("3.05" in line or "KH" in line for line in lines)


def test_strip_box_svg_handles_one_value_and_many() -> None:
    assert strip_box_svg([]) == ""
    one = strip_box_svg([2.5])
    assert one.count("<circle") == 1 and "<rect" in one
    many = strip_box_svg([1.0, 2.0, 3.0, 4.0, 10.0])
    assert many.count("<circle") == 5


# --- end to end ------------------------------------------------------------------------


def test_export_dataset_writes_outputs_and_guarded_share(tmp_path: Path) -> None:
    events = [ev(1, "typed", "-0.3")]
    store, resolved, derived = pipeline_objects(events)
    out, share = tmp_path / "outputs", tmp_path / "share"
    notices: list[str] = []
    files = export_dataset(
        out, share, store, resolved, derived, events, [], notice=notices.append
    )
    for name in (
        VALUES_LONG_FILE,
        DATASET_WIDE_FILE,
        STATUS_LONG_FILE,
        FREE_TEXT_FILE,
        DICTIONARY_FILE,
        WORKBOOK_FILE,
        QC_REPORT_FILE,
    ):
        assert (out / name).is_file(), name
    assert files.masked_cells == 0, "a writer leaked a token the guard had to mask"
    summary = (share / QC_SUMMARY_FILE).read_text(encoding="utf-8")
    assert MASK not in summary
    for secret in ("3.05", "11.pdf", "SN-1234", "Patient", "KH"):
        assert secret not in summary, secret
    var = read_rows(share / VARIABLE_SUMMARY_FILE)
    assert any(r["variable"] == "FEV1_pre" and r["n_read"] == "1" for r in var)
    assert "record_id" not in var[0]
    wb = load_workbook(out / WORKBOOK_FILE)
    assert wb.sheetnames == ["Data", "Status", "Dictionary", "Derived"]
    data = wb["Data"]
    header = [c.value for c in data[1]]
    assert header[:3] == ["record_id", "layout", "pages"]
    assert data.cell(row=2, column=header.index("FEV1_pre") + 1).value == 3.05
    wide = read_rows(out / DATASET_WIDE_FILE)
    assert wide[0]["FVC_z"] == "-0.3"
