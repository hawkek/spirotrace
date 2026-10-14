"""Coding workbooks: one per coder, blind, reproducible, nothing but ids and labels."""

import shutil
import subprocess
from pathlib import Path

import pytest
from conftest import design, synthetic_frame
from openpyxl import load_workbook

from spirotrace.evaluate import (
    CoreField,
    EvaluateError,
    allocate_coders,
    draw_sample,
    write_coder_workbooks,
)
from spirotrace.evaluate.workbooks import (
    DEFAULT_INSTRUCTIONS,
    FIRST_FIELD_ROW,
    HEADERS,
    INSTRUCTIONS_SHEET,
    VALUE_STATE_MEANINGS,
    check_core_fields,
    value_state_list,
)
from spirotrace.model import ColumnRole, Parameter, ValueState

CORE_FIELDS = {
    "layout_a": (
        CoreField(Parameter.FEV1, ColumnRole.PRE, "FEV1 pre"),
        CoreField(Parameter.FEV1, ColumnRole.POST, "FEV1 post"),
        CoreField(Parameter.AGE, ColumnRole.MEASURED, "Age"),
    ),
    "layout_b": (
        CoreField(Parameter.FEV1, ColumnRole.PRE, "FEV1 pre"),
        CoreField(Parameter.TLC, ColumnRole.MEASURED, "TLC"),
        CoreField(Parameter.TLC, ColumnRole.Z, "TLC z"),
        CoreField(Parameter.AGE, ColumnRole.MEASURED, "Age"),
    ),
}

FIXED_LABELS = {"Record", "Layout", "Start time", "End time", *HEADERS}


def cells(path: Path) -> dict[tuple[str, str], object]:
    """Every non-empty cell of a workbook by (sheet, coordinate)."""
    workbook = load_workbook(path)
    out: dict[tuple[str, str], object] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    out[(sheet.title, cell.coordinate)] = cell.value
    return out


@pytest.fixture
def allocated():
    frame = synthetic_frame()
    d = design(k=14, seed=4)
    sample = draw_sample(frame, d)
    return frame, d, sample, allocate_coders(sample, d, {})


def test_one_workbook_per_coder_with_their_reports_in_their_own_order(
    tmp_path, allocated
) -> None:
    _frame, d, _sample, allocation = allocated
    paths = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)
    assert set(paths) == set(d.coders)
    orders = {}
    for coder, path in paths.items():
        assert path.name == f"coding_{coder}_CONTAINS_DATA.xlsx"
        workbook = load_workbook(path)
        titles = workbook.sheetnames
        assert titles[0] == INSTRUCTIONS_SHEET
        expected = {a.record_id for a in allocation.for_coder(coder)}
        assert set(titles[1:]) == expected
        assert len(titles) == 15
        orders[coder] = titles[1:]
        assert workbook[INSTRUCTIONS_SHEET]["B1"].value == coder
        assert workbook[INSTRUCTIONS_SHEET]["B2"].value == 14
    assert orders["c1"] != sorted(orders["c1"])  # shuffled, not alphabetical
    overlap = next(a for a in allocation.assignments if a.overlap)
    first, second = overlap.coders
    assert orders[first].index(overlap.record_id) != orders[second].index(
        overlap.record_id
    )


def test_report_sheet_layout_dropdown_and_text_format(tmp_path, allocated) -> None:
    _frame, d, _sample, allocation = allocated
    paths = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)
    a = next(x for x in allocation.for_coder("c1") if x.layout == "layout_b")
    sheet = load_workbook(paths["c1"])[a.record_id]
    assert [sheet.cell(row=r, column=1).value for r in range(1, 5)] == [
        "Record",
        "Layout",
        "Start time",
        "End time",
    ]
    assert sheet["B1"].value == a.record_id
    assert sheet["B2"].value == "layout_b"
    assert sheet["B3"].value is None and sheet["B4"].value is None
    assert [c.value for c in sheet[FIRST_FIELD_ROW - 1]] == list(HEADERS)
    labels = [
        sheet.cell(row=r, column=1).value
        for r in range(FIRST_FIELD_ROW, FIRST_FIELD_ROW + 4)
    ]
    assert labels == ["FEV1 pre", "TLC", "TLC z", "Age"]
    assert sheet.cell(row=FIRST_FIELD_ROW + 4, column=1).value is None
    for r in range(FIRST_FIELD_ROW, FIRST_FIELD_ROW + 4):
        assert sheet.cell(row=r, column=2).number_format == "@"
        for col in range(2, 6):
            assert sheet.cell(row=r, column=col).value is None
    validations = sheet.data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].formula1 == '"value,empty,dash,not_on_page,unreadable"'
    assert validations[0].formula1 == value_state_list()
    assert str(validations[0].sqref) == f"C{FIRST_FIELD_ROW}:C{FIRST_FIELD_ROW + 3}"


def test_dropdown_and_meanings_cover_exactly_the_value_states() -> None:
    assert value_state_list() == '"' + ",".join(s.value for s in ValueState) + '"'
    assert set(VALUE_STATE_MEANINGS) == set(ValueState)


def test_every_cell_is_a_label_an_id_a_layout_or_the_instructions(
    tmp_path, allocated
) -> None:
    frame, d, _sample, allocation = allocated
    paths = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)
    field_labels = {f.label for fields in CORE_FIELDS.values() for f in fields}
    allowed_text = (
        FIXED_LABELS
        | field_labels
        | frame.record_ids
        | set(frame.layouts)
        | {"Coder", "Reports", "Value state", "Meaning", DEFAULT_INSTRUCTIONS}
        | {s.value for s in ValueState}
        | set(VALUE_STATE_MEANINGS.values())
        | set(d.coders)
    )
    for path in paths.values():
        for (sheet, _coordinate), value in cells(path).items():
            if sheet == INSTRUCTIONS_SHEET and value == 14:
                continue
            assert value in allowed_text, (sheet, value)


def test_sentinels_in_inputs_never_reach_a_workbook(tmp_path) -> None:
    # the only strings the generator receives besides ids and labels are feature names,
    # exclusion lists and coder labels; plant sentinels in the first two
    frame = synthetic_frame()
    sentinel_feature = "sentinel_feature_4711"
    from spirotrace.evaluate import FrameRecord, SamplingFrame

    spiked = SamplingFrame(
        tuple(
            FrameRecord(r.record_id, r.layout, r.features | {sentinel_feature})
            for r in frame.records
        )
    )
    d = design(k=14, seed=6)
    sample = draw_sample(spiked, d)
    allocation = allocate_coders(sample, d, {"c1": frozenset({"SENTINEL-NOT-AN-ID"})})
    paths = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)
    for path in paths.values():
        text = " ".join(str(v) for v in cells(path).values())
        assert sentinel_feature not in text
        assert "SENTINEL" not in text


def test_same_seed_gives_identical_cell_contents(tmp_path, allocated) -> None:
    _frame, d, _sample, allocation = allocated
    first = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path / "one", d.seed)
    second = write_coder_workbooks(allocation, CORE_FIELDS, tmp_path / "two", d.seed)
    for coder in d.coders:
        assert cells(first[coder]) == cells(second[coder])
        assert (
            load_workbook(first[coder]).sheetnames
            == load_workbook(second[coder]).sheetnames
        )
    other = write_coder_workbooks(
        allocation, CORE_FIELDS, tmp_path / "three", d.seed + 1
    )
    assert (
        load_workbook(first["c1"]).sheetnames != load_workbook(other["c1"]).sheetnames
    )


def test_refuses_to_overwrite(tmp_path, allocated) -> None:
    _frame, d, _sample, allocation = allocated
    write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)
    with pytest.raises(EvaluateError, match="refusing to overwrite"):
        write_coder_workbooks(allocation, CORE_FIELDS, tmp_path, d.seed)


def test_core_field_validation() -> None:
    with pytest.raises(EvaluateError, match="must be a Parameter"):
        CoreField("FEV1", ColumnRole.PRE, "x")  # type: ignore[arg-type]
    with pytest.raises(EvaluateError, match="must be a ColumnRole"):
        CoreField(Parameter.FEV1, "pre", "x")  # type: ignore[arg-type]
    with pytest.raises(EvaluateError, match="blank"):
        CoreField(Parameter.FEV1, ColumnRole.PRE, " ")
    with pytest.raises(EvaluateError, match="no core fields for layout 'layout_b'"):
        check_core_fields(
            {"layout_a": CORE_FIELDS["layout_a"]}, ("layout_a", "layout_b")
        )
    twice = (
        CORE_FIELDS["layout_a"][0],
        CoreField(Parameter.FEV1, ColumnRole.PRE, "again"),
    )
    with pytest.raises(EvaluateError, match="listed twice"):
        check_core_fields({"layout_a": twice}, ("layout_a",))
    same_label = (
        CORE_FIELDS["layout_a"][0],
        CoreField(Parameter.FVC, ColumnRole.PRE, "FEV1 pre"),
    )
    with pytest.raises(EvaluateError, match="label is used twice"):
        check_core_fields({"layout_a": same_label}, ("layout_a",))


@pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed")
def test_workbook_and_manifest_names_are_ignored_by_git(tmp_path: Path) -> None:
    repo_root = Path(__file__).resolve().parent.parent
    shutil.copyfile(repo_root / ".gitignore", tmp_path / ".gitignore")
    subprocess.run(["git", "init", "-q"], cwd=tmp_path, check=True)
    for name in ("coding_c1_CONTAINS_DATA.xlsx", "sample_manifest_CONTAINS_DATA.csv"):
        for rel in (f"anywhere/{name}", name):
            result = subprocess.run(
                ["git", "check-ignore", "-q", rel], cwd=tmp_path, check=False
            )
            assert result.returncode == 0, rel
