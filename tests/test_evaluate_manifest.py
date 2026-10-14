"""Manifest, counts summary and the one-call generator. Console output holds no record id."""

import csv
from pathlib import Path

import pytest
from conftest import FEATURES, design, synthetic_frame
from openpyxl import load_workbook
from test_evaluate_workbooks import CORE_FIELDS

from spirotrace.evaluate import (
    COUNTS_NAME,
    MANIFEST_NAME,
    EvaluateError,
    allocate_coders,
    count_summary,
    draw_sample,
    generate_coding_sheets,
    write_manifest,
)
from spirotrace.evaluate.manifest import (
    COUNT_COLUMNS,
    MANIFEST_COLUMNS,
    format_counts,
    write_counts,
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


@pytest.fixture
def allocated():
    frame = synthetic_frame()
    d = design(k=14, seed=8)
    sample = draw_sample(frame, d)
    return frame, d, sample, allocate_coders(sample, d, {})


def test_manifest_rows(tmp_path, allocated) -> None:
    _frame, _d, sample, allocation = allocated
    path = write_manifest(allocation, sample, tmp_path / MANIFEST_NAME)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = read_csv(path)
    assert list(rows[0]) == list(MANIFEST_COLUMNS)
    assert len(rows) == 36
    assert [r["record_id"] for r in rows] == sorted(sample.record_ids)
    overlap = [r for r in rows if r["overlap"] == "true"]
    assert len(overlap) == 6
    assert all(len(r["coders"].split()) == 2 for r in overlap)
    assert all(len(r["coders"].split()) == 1 for r in rows if r["overlap"] == "false")
    for row in rows:
        record = sample[row["record_id"]]
        assert row["layout"] == record.layout
        assert row["stratum"] == record.stratum
        assert int(row["stratum_size"]) == record.stratum_size
        assert int(row["sampled_in_stratum"]) == record.sampled_in_stratum
        assert float(row["sampling_fraction"]) == record.sampling_fraction
        assert row["seed"] == "8"
        assert tuple(row["coders"].split()) == allocation[row["record_id"]].coders
    with pytest.raises(EvaluateError, match="refusing to overwrite"):
        write_manifest(allocation, sample, path)


def test_count_summary_has_no_record_id_and_adds_up(allocated) -> None:
    frame, _d, sample, allocation = allocated
    rows = count_summary(allocation, sample)
    kinds = [r.kind for r in rows]
    assert kinds == ["stratum"] * len(sample.strata) + ["layout", "layout"] + [
        "coder"
    ] * 3 + ["total"]
    strata = [r for r in rows if r.kind == "stratum"]
    assert sum(r.sampled for r in strata) == 36
    assert sum(r.frame_size or 0 for r in strata) == len(frame)
    assert sum(r.overlap for r in strata) == 6
    layouts = {r.name: r for r in rows if r.kind == "layout"}
    assert layouts["layout_a"].frame_size == 58 and layouts["layout_b"].frame_size == 44
    assert layouts["layout_a"].sampled + layouts["layout_b"].sampled == 36
    assert layouts["layout_a"].overlap == layouts["layout_b"].overlap == 3
    coders = {r.name: r for r in rows if r.kind == "coder"}
    assert all(
        r.sampled == 14 and r.overlap == 4 and r.frame_size is None
        for r in coders.values()
    )
    total = rows[-1]
    assert (total.frame_size, total.sampled, total.overlap) == (len(frame), 36, 6)
    assert total.sampling_fraction == pytest.approx(36 / len(frame))
    text = format_counts(rows)
    assert not any(rid in text for rid in frame.record_ids)
    assert "layout_a:hard" in text and "c3" in text


def test_counts_file(tmp_path, allocated) -> None:
    frame, _d, sample, allocation = allocated
    rows = count_summary(allocation, sample)
    path = write_counts(rows, tmp_path / COUNTS_NAME)
    written = read_csv(path)
    assert list(written[0]) == list(COUNT_COLUMNS)
    assert len(written) == len(rows)
    assert written[-1]["kind"] == "total"
    coder_row = next(r for r in written if r["kind"] == "coder")
    assert coder_row["frame_size"] == "" and coder_row["sampling_fraction"] == ""
    assert not any(
        rid in path.read_text(encoding="utf-8-sig") for rid in frame.record_ids
    )
    with pytest.raises(EvaluateError, match="refusing to overwrite"):
        write_counts(rows, path)


def test_generate_writes_everything_and_prints_counts_only(tmp_path, capsys) -> None:
    frame = synthetic_frame()
    d = design(k=(14, 14, 10), seed=21)
    exclusions = {"c1": frozenset(sorted(frame.record_ids)[:30])}
    out = generate_coding_sheets(frame, d, exclusions, CORE_FIELDS, tmp_path / "coding")
    assert out.manifest.name == MANIFEST_NAME and out.manifest.exists()
    assert out.counts.name == COUNTS_NAME and out.counts.exists()
    assert set(out.workbooks) == {"c1", "c2", "c3"}
    assert all(p.exists() for p in out.workbooks.values())
    assert len(out.sample) == len(out.allocation) == 32
    assert not {a.record_id for a in out.allocation.for_coder("c1")} & exclusions["c1"]
    assert len(load_workbook(out.workbooks["c3"]).sheetnames) == 11
    printed = capsys.readouterr().out
    assert not any(rid in printed for rid in frame.record_ids)
    assert "seed 21: 32 reports, 3 workbooks written" in printed
    assert "layout_b:volumes" in printed
    # a second run into the same folder writes nothing
    with pytest.raises(EvaluateError, match="refusing to overwrite"):
        generate_coding_sheets(frame, d, exclusions, CORE_FIELDS, tmp_path / "coding")


def test_generate_checks_before_writing(tmp_path) -> None:
    frame = synthetic_frame()
    d = design(k=14, seed=1)
    with pytest.raises(EvaluateError, match="no core fields"):
        generate_coding_sheets(
            frame, d, {}, {"layout_a": CORE_FIELDS["layout_a"]}, tmp_path
        )
    with pytest.raises(EvaluateError, match="eligible records in the frame"):
        generate_coding_sheets(
            frame, d, {"c1": frame.record_ids}, CORE_FIELDS, tmp_path
        )
    assert list(tmp_path.iterdir()) == []
    assert FEATURES == d.precedence
